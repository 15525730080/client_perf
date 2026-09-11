# coding: utf-8
"""
export_service.py — Excel 导出服务层（集中自 api.py）。

保留原函数名：create_excel_report / create_comparison_excel /
export_label_comparison_excel_func，并新增 export_advanced_excel。
所有函数返回绝对路径，并支持通过 output_dir 注入输出目录（便于测试）。

解决的问题：
  * 工作表名称非法字符（\\ / * ? : [ ]）、重复名、31 字符限制
  * 文件名非法字符清理
"""
from __future__ import annotations

import re
import time
from datetime import datetime
from pathlib import Path
from typing import Any

from openpyxl import Workbook

from client_perf.log import log as logger
from client_perf.paths import get_data_dir

# 默认写入用户可写运行数据目录，可通过 CLIENT_PERF_DATA_DIR 覆盖。
_DATA_DIR = get_data_dir()
REPORT_DIR = _DATA_DIR / "test_result" / "excel_reports"
COMPARISON_REPORT_DIR = _DATA_DIR / "test_result" / "comparison_reports"

# ── 名称清理 ────────────────────────────────────────────────────

_ILLEGAL_SHEET_CHARS = re.compile(r"[\\/*?:\[\]]")
_ILLEGAL_FILENAME_CHARS = re.compile(r'[\\/*?:"<>|]')
_MAX_SHEET_NAME_LEN = 31


def _sanitize_sheet_name(name: str, used: set[str]) -> str:
    """清理工作表名称：去除非法字符、限制 31 字符、保证唯一。"""
    if not name:
        name = "Sheet"
    name = _ILLEGAL_SHEET_CHARS.sub("_", name)
    if len(name) > _MAX_SHEET_NAME_LEN:
        name = name[:_MAX_SHEET_NAME_LEN]

    base = name
    suffix = 1
    while name in used:
        suffix_str = f"_{suffix}"
        if len(base) + len(suffix_str) > _MAX_SHEET_NAME_LEN:
            base = base[: _MAX_SHEET_NAME_LEN - len(suffix_str)]
        name = base + suffix_str
        suffix += 1
    used.add(name)
    return name


def _sanitize_filename(name: str) -> str:
    """清理文件名中的非法字符。"""
    if not name:
        name = "report"
    name = _ILLEGAL_FILENAME_CHARS.sub("_", name)
    name = name.strip().rstrip(".")
    if len(name) > 200:
        name = name[:200]
    return name


def _resolve_output_dir(output_dir: str | Path | None, default: Path) -> Path:
    out = Path(output_dir) if output_dir else default
    out.mkdir(parents=True, exist_ok=True)
    return out


# 指标展示顺序（与历史一致）
_METRIC_COLUMNS = [
    ("cpu", "CPU 使用率 (%)"),
    ("memory", "内存使用量 (MB)"),
    ("fps", "FPS"),
    ("gpu", "GPU 使用率 (%)"),
    ("threads", "线程数"),
    ("handles", "句柄数"),
    ("disk_read", "磁盘读取 (MB/s)"),
    ("disk_write", "磁盘写入 (MB/s)"),
    ("net_sent", "网络发送 (MB/s)"),
    ("net_recv", "网络接收 (MB/s)"),
]

_AVG_HEADERS = [
    "任务名称", "版本", "基准",
    "CPU 均值(%)", "内存均值(MB)", "FPS 均值", "GPU 均值(%)",
    "线程数均值", "句柄数均值",
    "磁盘读取均值(MB/s)", "磁盘写入均值(MB/s)",
    "网络发送均值(MB/s)", "网络接收均值(MB/s)",
]


# ── 1. 单任务 Excel 报告 ────────────────────────────────────────

def create_excel_report(
    task_name: str,
    data: list[dict],
    save_dir: str | None = None,
    output_dir: str | Path | None = None,
) -> str:
    """导出单个任务的 Excel 报告，返回绝对路径。"""
    try:
        wb = Workbook()
        used: set[str] = set()

        for item in data:
            sheet_name = _sanitize_sheet_name(item.get("name", "未知"), used)
            ws = wb.create_sheet(title=sheet_name)

            values = item.get("value", [])
            if values:
                all_fields: set = set()
                for row in values:
                    all_fields.update(row.keys())
                headers = sorted(all_fields)

                ws.append(headers)
                for row in values:
                    ws.append([row.get(h) for h in headers])
            else:
                ws.append(["无数据"])

        if "Sheet" in wb.sheetnames:
            wb.remove(wb["Sheet"])

        timestamp = time.strftime("%Y%m%d_%H%M%S")
        filename = _sanitize_filename(f"{task_name}_{timestamp}.xlsx")
        out_dir = _resolve_output_dir(output_dir, REPORT_DIR)
        file_path = out_dir / filename

        wb.save(file_path)
        logger.info(f"Excel 报告导出成功: {file_path}")
        return str(Path(file_path).resolve())
    except Exception as e:
        logger.error(f"导出 Excel 报告失败: {e}")
        raise


# ── 2. 任务对比 Excel ──────────────────────────────────────────

async def create_comparison_excel(
    data: dict,
    report_name: str,
    output_dir: str | Path | None = None,
) -> str:
    """导出任务对比报告为 Excel，返回绝对路径。"""
    try:
        wb = Workbook()
        used: set[str] = set()

        ws = wb.active
        ws.title = _sanitize_sheet_name("对比结果", used)
        ws.append(_AVG_HEADERS)

        for task in data.get("tasks", []):
            avg = task.get("avg", {})
            ws.append([
                task.get("name", ""),
                task.get("version", ""),
                "是" if task.get("id") == data.get("base_task", {}).get("id") else "否",
                avg.get("cpu_avg", None),
                avg.get("memory_avg", None),
                avg.get("fps_avg", None),
                avg.get("gpu_avg", None),
                avg.get("threads_avg", None),
                avg.get("handles_avg", None),
                avg.get("disk_read_avg", None),
                avg.get("disk_write_avg", None),
                avg.get("net_sent_avg", None),
                avg.get("net_recv_avg", None),
            ])

        # 每个任务的原始数据工作表
        for task in data.get("tasks", []):
            task_name = task.get("name", "未知任务")
            task_id = task.get("id", "")
            sheet_name = _sanitize_sheet_name(f"{task_name}_{task_id}", used)

            task_ws = wb.create_sheet(title=sheet_name)
            task_data = task.get("data", {})

            row_idx = 1
            for metric_key, metric_name in _METRIC_COLUMNS:
                metric_data = task_data.get(metric_key, [])
                if not metric_data:
                    continue

                task_ws.cell(row=row_idx, column=1, value=metric_name)
                row_idx += 1
                task_ws.cell(row=row_idx, column=1, value="时间")
                task_ws.cell(row=row_idx, column=2, value="值")
                row_idx += 1

                for item in metric_data:
                    time_val = item.get("time")
                    value_val = item.get("value")
                    if time_val is not None:
                        time_str = datetime.fromtimestamp(time_val).strftime("%Y-%m-%d %H:%M:%S")
                        task_ws.cell(row=row_idx, column=1, value=time_str)
                        task_ws.cell(row=row_idx, column=2, value=value_val)
                        row_idx += 1

                row_idx += 1  # 空行分隔

        timestamp = time.strftime("%Y%m%d_%H%M%S")
        filename = _sanitize_filename(f"{report_name}_{timestamp}.xlsx")
        out_dir = _resolve_output_dir(output_dir, COMPARISON_REPORT_DIR)
        file_path = out_dir / filename

        wb.save(file_path)
        logger.info(f"对比报告导出成功: {file_path}")
        return str(Path(file_path).resolve())
    except Exception as e:
        logger.error(f"导出对比报告失败: {e}")
        raise


# ── 3. 标签对比 Excel ──────────────────────────────────────────

async def export_label_comparison_excel_func(
    data: dict,
    report_name: str,
    output_dir: str | Path | None = None,
) -> str:
    """导出标签对比报告为 Excel，返回绝对路径。

    task 的 data 由 extract_detailed_data 生成，包含：
      - "_aligned": { metric: [v, ...] } 与 timestamps 对齐的标量数组
      - "timestamps": 对齐参考时间轴
    """
    try:
        wb = Workbook()
        used: set[str] = set()

        ws = wb.active
        ws.title = _sanitize_sheet_name("标签对比结果", used)
        headers = [
            "任务名称", "版本", "标签名称", "开始时间", "结束时间",
            "CPU 均值(%)", "内存均值(MB)", "FPS 均值", "GPU 均值(%)",
            "线程数均值", "句柄数均值", "磁盘读取均值(MB/s)",
            "磁盘写入均值(MB/s)", "网络发送均值(MB/s)", "网络接收均值(MB/s)",
        ]
        ws.append(headers)

        for task in data.get("tasks", []):
            avg = task.get("avg", {})
            ws.append([
                task.get("name", ""),
                task.get("version", ""),
                task.get("label_name", ""),
                datetime.fromtimestamp(task.get("start_ts", 0)).strftime("%Y-%m-%d %H:%M:%S"),
                datetime.fromtimestamp(task.get("end_ts", 0)).strftime("%Y-%m-%d %H:%M:%S"),
                avg.get("cpu_avg", None),
                avg.get("memory_avg", None),
                avg.get("fps_avg", None),
                avg.get("gpu_avg", None),
                avg.get("threads_avg", None),
                avg.get("handles_avg", None),
                avg.get("disk_read_avg", None),
                avg.get("disk_write_avg", None),
                avg.get("net_sent_avg", None),
                avg.get("net_recv_avg", None),
            ])

        # 每个标签的原始数据工作表（使用对齐数组）
        for task in data.get("tasks", []):
            task_name = task.get("name", "未知任务")
            label_name = task.get("label_name", "未知标签")
            sheet_name = _sanitize_sheet_name(f"{task_name}_{label_name}", used)

            task_ws = wb.create_sheet(title=sheet_name)
            task_data = task.get("data", {})
            aligned_map = task_data.get("_aligned", {})
            timestamps = task_data.get("timestamps", [])

            row_idx = 1
            for metric_key, metric_name in _METRIC_COLUMNS:
                metric_data = aligned_map.get(metric_key)
                if metric_data is None:
                    metric_data = task_data.get(metric_key, [])

                if not metric_data or not timestamps:
                    continue

                task_ws.cell(row=row_idx, column=1, value=metric_name)
                row_idx += 1
                task_ws.cell(row=row_idx, column=1, value="时间")
                task_ws.cell(row=row_idx, column=2, value="值")
                row_idx += 1

                for i, value in enumerate(metric_data):
                    if i < len(timestamps) and value is not None:
                        time_val = timestamps[i]
                        time_str = datetime.fromtimestamp(time_val).strftime("%Y-%m-%d %H:%M:%S")
                        task_ws.cell(row=row_idx, column=1, value=time_str)
                        task_ws.cell(row=row_idx, column=2, value=value)
                        row_idx += 1

                row_idx += 1  # 空行分隔

        timestamp = time.strftime("%Y%m%d_%H%M%S")
        filename = _sanitize_filename(f"{report_name}_{timestamp}.xlsx")
        out_dir = _resolve_output_dir(output_dir, COMPARISON_REPORT_DIR)
        file_path = out_dir / filename

        wb.save(file_path)
        logger.info(f"标签对比报告导出成功: {file_path}")
        return str(Path(file_path).resolve())
    except Exception as e:
        logger.error(f"导出标签对比报告失败: {e}")
        raise


# ── 4. 高级对比 Excel ──────────────────────────────────────────

async def export_advanced_excel(
    data: dict,
    report_name: str,
    output_dir: str | Path | None = None,
) -> str:
    """导出高级对比（统计显著性 + 异常值 + 瓶颈）为 Excel，返回绝对路径。"""
    try:
        wb = Workbook()
        used: set[str] = set()

        base_task = data.get("base_task", {})
        cmp_task = data.get("compare_task", {})
        analysis = data.get("advanced_analysis", {})
        significance = analysis.get("statistical_significance", {})
        outliers = analysis.get("outliers", {})
        bottleneck = analysis.get("bottleneck_analysis", {})

        # 概览
        ws = wb.active
        ws.title = _sanitize_sheet_name("高级对比概览", used)
        ws.append(["基准任务", base_task.get("name", "")])
        ws.append(["对比任务", cmp_task.get("name", "")])
        ws.append(["瓶颈_summary", bottleneck.get("summary", "")])
        ws.append([])

        # 显著性分析
        sig_ws = wb.create_sheet(title=_sanitize_sheet_name("显著性分析", used))
        sig_ws.append([
            "指标", "基准均值", "对比均值", "差值", "变化百分比(%)",
            "p_value", "置信度(%)", "是否显著",
        ])
        for metric, s in significance.items():
            sig_ws.append([
                metric,
                s.get("base_mean"),
                s.get("compare_mean"),
                s.get("diff"),
                s.get("percent_change"),
                s.get("p_value"),
                s.get("confidence"),
                "是" if s.get("is_significant") else "否",
            ])

        # 异常值检测
        out_ws = wb.create_sheet(title=_sanitize_sheet_name("异常值检测", used))
        out_ws.append(["指标", "异常值", "Z分数", "是否偏高(>均值)"])
        for metric, o in outliers.items():
            for item in o.get("outliers", []):
                out_ws.append([
                    metric,
                    item.get("value"),
                    item.get("z_score"),
                    "是" if item.get("is_high") else "否",
                ])

        # 瓶颈分析
        bn_ws = wb.create_sheet(title=_sanitize_sheet_name("瓶颈分析", used))
        bn_ws.append(["指标", "变化百分比(%)", "是否变差"])
        for b in bottleneck.get("potential_bottlenecks", []):
            bn_ws.append([
                b.get("metric"),
                b.get("percent_change"),
                "是" if b.get("is_worse") else "否",
            ])

        timestamp = time.strftime("%Y%m%d_%H%M%S")
        filename = _sanitize_filename(f"{report_name}_{timestamp}.xlsx")
        out_dir = _resolve_output_dir(output_dir, COMPARISON_REPORT_DIR)
        file_path = out_dir / filename

        wb.save(file_path)
        logger.info(f"高级对比报告导出成功: {file_path}")
        return str(Path(file_path).resolve())
    except Exception as e:
        logger.error(f"导出高级对比报告失败: {e}")
        raise
